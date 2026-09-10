"""Section 2 — DOCX text extraction (python-docx).

Paragraphs first, then table cells (tables scramble chronology in real ATS —
we at least concatenate them last so the main flow stays readable).

API:
    extract_docx(path_or_stream) -> DocxExtraction(text, warnings)
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

try:
    import docx as python_docx
except ImportError:  # pragma: no cover
    python_docx = None


@dataclass
class DocxExtraction:
    text: str = ""
    warnings: list[str] = field(default_factory=list)


def extract_docx(source) -> DocxExtraction:
    if python_docx is None:
        return DocxExtraction(warnings=["python-docx not installed — run: pip install python-docx"])
    result = DocxExtraction()
    try:
        stream = io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
        document = python_docx.Document(stream)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        result.text = "\n".join(parts)
        if not result.text.strip():
            result.warnings.append("DOCX contained no extractable text")
    except Exception as exc:  # noqa: BLE001
        result.warnings.append(f"DOCX extraction failed: {type(exc).__name__}: {exc}")
    return result


if __name__ == "__main__":
    import sys

    if len(sys.argv) > 1:
        extraction = extract_docx(sys.argv[1])
        print(extraction.text[:1500])
        print("warnings:", extraction.warnings)
    else:
        print("usage: python docx_extractor.py <file.docx>")
